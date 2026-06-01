from __future__ import annotations

import io
import os
import re
from dataclasses import dataclass


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    text: str


def _normalize_text(text: str) -> str:
    if not text:
        return ""

    # Normalize whitespace/newlines.
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = t.replace("\u00a0", " ")  # NBSP
    t = t.replace("\u200b", "")  # zero-width space
    t = re.sub(r"[\t\v\f]+", " ", t)

    lines = []
    for raw_line in t.split("\n"):
        line = raw_line.strip()

        # Drop common PDF artifacts: bullet-only lines and isolated page numbers.
        if line in {"•", "·", "●", "◦", "-", "–", "—"}:
            continue
        if re.fullmatch(r"\d{1,3}", line):
            continue

        # Collapse runs of bullets into a single hyphen.
        line = re.sub(r"^[•·●◦]+\s*", "- ", line)

        lines.append(line)

    # Remove excessive blank lines.
    out_lines: list[str] = []
    blank_run = 0
    for line in lines:
        if not line:
            blank_run += 1
            if blank_run <= 2:
                out_lines.append("")
            continue
        blank_run = 0
        out_lines.append(line)

    return "\n".join(out_lines).strip()


def _read_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _read_pdf_with_pymupdf(file_bytes: bytes) -> str:
    import fitz  # PyMuPDF

    text_parts: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts).strip()


def _read_pdf_with_pypdf2(file_bytes: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def _read_docx(file_bytes: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs).strip()


def extract_text(filename: str, file_bytes: bytes) -> ExtractedDocument:
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".txt":
        return ExtractedDocument(filename=filename, text=_normalize_text(_read_txt(file_bytes)))

    if ext == ".docx":
        return ExtractedDocument(filename=filename, text=_normalize_text(_read_docx(file_bytes)))

    if ext == ".pdf":
        # Prefer PyMuPDF (usually best extraction), fallback to PyPDF2.
        try:
            text = _read_pdf_with_pymupdf(file_bytes)
        except Exception:
            text = _read_pdf_with_pypdf2(file_bytes)

        return ExtractedDocument(filename=filename, text=_normalize_text(text))

    raise ValueError(f"Unsupported file type: {ext}. Upload .pdf, .docx, or .txt")
